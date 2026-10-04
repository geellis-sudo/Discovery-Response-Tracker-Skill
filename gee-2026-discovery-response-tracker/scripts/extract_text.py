#!/usr/bin/env python3
"""
extract_text.py — pull plain text out of a discovery document, whatever
format it happens to be in.

Handles:
  - .docx           via python-docx
  - .pdf (text)      via pdfplumber
  - .pdf (scanned)   OCR fallback via pdf2image + pytesseract, page by page,
                     only for pages where the text layer comes back empty
  - .txt / .md       passthrough

Inserts "--- page N ---" markers for PDFs so downstream extraction can cite
page numbers in source_locator.

Usage:
    python extract_text.py <input_file> [--out <output_file>]

If --out is omitted, prints to stdout.
"""
import argparse
import sys
from pathlib import Path

MIN_CHARS_PER_PAGE = 20  # below this, treat a PDF page as "no text layer" and OCR it


def extract_docx(path: Path) -> str:
    try:
        import docx
    except ImportError:
        raise SystemExit(
            "Missing dependency 'python-docx'. Install with: pip install python-docx --break-system-packages"
        )
    d = docx.Document(str(path))
    parts = []
    for para in d.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def ocr_pdf_page(path: Path, page_index: int) -> str:
    try:
        from pdf2image import convert_from_path
        import pytesseract
    except ImportError:
        raise SystemExit(
            "Missing OCR dependencies. Install with: "
            "pip install pdf2image pytesseract --break-system-packages "
            "(also requires the system packages 'poppler' and 'tesseract-ocr')."
        )
    images = convert_from_path(
        str(path), first_page=page_index + 1, last_page=page_index + 1
    )
    if not images:
        return ""
    return pytesseract.image_to_string(images[0])


def extract_pdf(path: Path) -> str:
    try:
        import pdfplumber
    except ImportError:
        raise SystemExit(
            "Missing dependency 'pdfplumber'. Install with: pip install pdfplumber --break-system-packages"
        )
    parts = []
    with pdfplumber.open(str(path)) as pdf:
        for i, page in enumerate(pdf.pages):
            text = (page.extract_text() or "").strip()
            if len(text) < MIN_CHARS_PER_PAGE:
                ocr_text = ocr_pdf_page(path, i).strip()
                text = ocr_text if ocr_text else text
            parts.append(f"--- page {i + 1} ---\n{text}")
    return "\n\n".join(parts)


def extract_plain(path: Path) -> str:
    return path.read_text(errors="replace")


def extract(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix in (".txt", ".md"):
        return extract_plain(path)
    raise SystemExit(
        f"Unsupported file type: {suffix}. Supported: .docx, .pdf, .txt, .md"
    )


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_file", type=Path)
    parser.add_argument("--out", type=Path, default=None)
    args = parser.parse_args()

    if not args.input_file.exists():
        raise SystemExit(f"File not found: {args.input_file}")

    text = extract(args.input_file)

    if args.out:
        args.out.write_text(text)
        print(f"Wrote {len(text)} chars to {args.out}", file=sys.stderr)
    else:
        print(text)


if __name__ == "__main__":
    main()
